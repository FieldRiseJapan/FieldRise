import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

sys.path.insert(0, str(Path(__file__).parent / "src"))
from pipeline import WireRow, classify_header, write_excel


def check(condition, label):
    if not condition:
        raise AssertionError(label)
    print(f"OK  {label}")


with tempfile.TemporaryDirectory() as temp_dir:
    output_dir = Path(temp_dir)
    header_crop = output_dir / "P1_F1_header.png"
    Image.new("RGB", (20, 10), "white").save(header_crop)
    rows = [
        WireRow(1, "P1-F1", "TF", "マルチコネクター", "1", "FL1A11", "", "R-Y2", ["Y2"], "未確認", "", "", str(header_crop), "5A", header_crop=str(header_crop)),
        WireRow(1, "P1-F1", "TF", "マルチコネクター", "2", "FL1A21", "", "R-Y5", ["Y5"], "未確認", "", "", str(header_crop), "5A", header_crop=str(header_crop)),
        WireRow(1, "P1-F2", "IF", "コネクター", "1", "IF1", "", "R-Y2", ["Y2"], "未確認", "", "", str(header_crop), "5A", header_crop=str(header_crop)),
    ]
    check(classify_header("IF") == "コネクター", "IF is classified as connector")
    output = write_excel(rows, Path("many_pages.pdf"), "ORDER123", output_dir, lambda *_: None, page_panels={1: "5A"})
    workbook = load_workbook(output, data_only=True)
    sheet = workbook["2"]
    check([cell.value for cell in sheet[5]][:3] == ["マーク主文字", "マーク個数", "解析スクリーンショット"], "screenshot is third output column")
    check(len(getattr(sheet, "_images", [])) == 2, "same header image is attached to both size rows")
    check(workbook["5"].cell(8, 3).value is None, "different wire size keeps screenshot column layout")
    check(output.exists(), "multi-page-style workbook is written")

print("MULTIPAGE AND HEADER IMAGE TESTS PASSED")
