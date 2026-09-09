from pathlib import Path
from openpyxl import load_workbook

path = Path('/home/ubuntu/25JNG38201W.xlsx')
wb = load_workbook(path, data_only=True)
print('SHEETS=', ','.join(wb.sheetnames))
for name in wb.sheetnames:
    ws = wb[name]
    headers = [ws.cell(5, col).value for col in range(1, 15)]
    print(f'{name}: headers={headers!r} images={len(getattr(ws, "_images", []))} rows={ws.max_row}')
print('IF_KIND=', 'コネクター')
